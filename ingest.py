#!/usr/bin/env python3
"""SAGE ingestion pipeline with schematic-aware PDF extraction improvements."""

import argparse
import csv
import hashlib
import json
import os
import re
import zipfile
from io import BytesIO
from typing import Dict, List

import chromadb
import docx
import fitz  # PyMuPDF
import openpyxl
from chromadb.utils.embedding_functions import SentenceTransformerEmbeddingFunction

from config import *

try:
    import pytesseract
    from PIL import Image

    OCR_AVAILABLE = True
except Exception:
    OCR_AVAILABLE = False


# -----------------------------
# Chunking
# -----------------------------
def _split_recursive(text: str, chunk_size: int, separators: List[str]) -> List[str]:
    text = text.strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]
    if not separators:
        return [text[i : i + chunk_size] for i in range(0, len(text), chunk_size)]

    sep = separators[0]
    if sep in text:
        raw_parts = text.split(sep)
        parts = []
        current = ""
        for part in raw_parts:
            candidate = part if not current else f"{current}{sep}{part}"
            if len(candidate) <= chunk_size:
                current = candidate
            else:
                if current:
                    parts.append(current)
                current = part
        if current:
            parts.append(current)

        if len(parts) > 1:
            out = []
            for p in parts:
                out.extend(_split_recursive(p, chunk_size, separators[1:]))
            return out

    return _split_recursive(text, chunk_size, separators[1:])


def chunk_text_semantic(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Split text into semantic-ish chunks with overlap."""
    separators = ["\n\n", "\n", ". ", "! ", "? ", "; ", ", ", " "]
    base_chunks = _split_recursive(text, chunk_size, separators)
    if not base_chunks:
        return []

    chunks = [base_chunks[0]]
    for i in range(1, len(base_chunks)):
        tail = chunks[-1][-overlap:] if overlap > 0 else ""
        chunks.append((tail + base_chunks[i]).strip())
    return chunks


def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Backwards-compatible alias."""
    return chunk_text_semantic(text, chunk_size=chunk_size, overlap=overlap)


# -----------------------------
# Schematic PDF extraction helpers
# -----------------------------
def _merge_spaced_letters(line: str) -> str:
    # Convert patterns like: "V e r t i c a l" -> "Vertical"
    def _collapse(match: re.Match) -> str:
        return match.group(0).replace(" ", "")

    return re.sub(r"\b(?:[A-Za-z]\s+){2,}[A-Za-z]\b", _collapse, line)


def _normalize_identifier_spacing(line: str) -> str:
    # Preserve/repair technical identifiers with accidental split: "J 4" -> "J4"
    line = re.sub(r"\b([A-Za-z]{1,4})\s+([0-9]{1,4}[A-Za-z]?)\b", r"\1\2", line)
    # Common terminal/channel forms: "X 1 : 3" -> "X1:3"
    line = re.sub(r"\b([A-Za-z]{1,4})([0-9]{1,4})\s*:\s*([0-9]{1,4})\b", r"\1\2:\3", line)
    return line


def _is_grid_coordinate_line(line: str) -> bool:
    # Drop likely CAD border/grid labels, but keep known component patterns like J4/P1/M3
    s = line.strip()
    if not s:
        return True

    # Keep explicit technical IDs
    if re.fullmatch(r"[A-Za-z]{1,3}[0-9]{1,4}[A-Za-z]?", s):
        return False

    # Likely coordinate artifacts
    if re.fullmatch(r"[A-Za-z]", s):
        return True
    if re.fullmatch(r"[0-9]{1,2}", s):
        return True
    if re.fullmatch(r"[A-Za-z]-?[0-9]{1,2}", s):
        return True

    return False


def _is_border_artifact(line: str) -> bool:
    s = line.strip()
    return bool(re.fullmatch(r"[-_=|.]{3,}", s))


def _is_noisy_text(raw_text: str) -> bool:
    lines = [ln.strip() for ln in raw_text.splitlines() if ln.strip()]
    if len(lines) < MIN_LINES_FOR_NOISE_CHECK:
        return False
    single_char = sum(1 for ln in lines if len(ln) == 1)
    ratio = single_char / max(1, len(lines))
    return ratio >= NOISY_SINGLE_CHAR_RATIO


def _preprocess_schematic_text(raw_text: str) -> str:
    cleaned = []
    for line in raw_text.splitlines():
        s = line.strip()
        if not s:
            continue
        if _is_border_artifact(s):
            continue

        s = _merge_spaced_letters(s)
        s = _normalize_identifier_spacing(s)

        if _is_grid_coordinate_line(s):
            continue

        s = re.sub(r"\s+", " ", s).strip()
        if s:
            cleaned.append(s)

    return "\n".join(cleaned)


def _dedupe_lines_keep_order(text: str) -> str:
    seen = set()
    out = []
    for line in text.splitlines():
        key = line.strip()
        if key and key not in seen:
            seen.add(key)
            out.append(key)
    return "\n".join(out)


def _try_ocr_page(page: fitz.Page) -> str:
    if not OCR_AVAILABLE:
        return ""
    try:
        mat = fitz.Matrix(300 / 72, 300 / 72)  # ~300 DPI
        pix = page.get_pixmap(matrix=mat, alpha=False)
        img = Image.open(BytesIO(pix.tobytes("png")))
        text = pytesseract.image_to_string(img)
        return text or ""
    except Exception:
        return ""


# -----------------------------
# Loaders
# -----------------------------
def load_pdf(filepath: str) -> List[Dict]:
    """Extract per-page PDF text with schematic-aware cleanup and OCR fallback."""
    entries: List[Dict] = []
    with fitz.open(filepath) as doc:
        for idx, page in enumerate(doc, start=1):
            native_raw = page.get_text("text", sort=True)  # layout-aware order for CAD drawings
            native_text = _preprocess_schematic_text(native_raw)

            trigger_ocr = (len(native_text) < OCR_TEXT_THRESHOLD) or _is_noisy_text(native_raw)
            ocr_text = _preprocess_schematic_text(_try_ocr_page(page)) if trigger_ocr else ""

            merged = "\n".join([t for t in [native_text, ocr_text] if t.strip()]).strip()
            merged = _dedupe_lines_keep_order(merged)

            if merged:
                entries.append(
                    {
                        "text": merged,
                        "source": filepath,
                        "page": idx,
                        "doc_type": "pdf",
                    }
                )
    return entries


def load_word(filepath: str) -> List[Dict]:
    doc = docx.Document(filepath)
    text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    return [{"text": text, "source": filepath, "doc_type": "docx"}] if text.strip() else []


def load_excel(filepath: str) -> List[Dict]:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join([str(cell) for cell in row if cell is not None])
            if row_text.strip():
                rows.append(row_text)
    text = "\n".join(rows)
    return [{"text": text, "source": filepath, "doc_type": "xlsx"}] if text.strip() else []


def load_text(filepath: str) -> List[Dict]:
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return [{"text": text, "source": filepath, "doc_type": "txt"}] if text.strip() else []


def load_csv(filepath: str) -> List[Dict]:
    lines = []
    with open(filepath, newline="", encoding="utf-8", errors="ignore") as f:
        reader = csv.DictReader(f)
        for row in reader:
            lines.append(" | ".join(f"{k}: {v}" for k, v in row.items() if v))
    text = "\n".join([ln for ln in lines if ln.strip()])
    return [{"text": text, "source": filepath, "doc_type": "csv"}] if text.strip() else []


def load_code(filepath: str, language: str) -> List[Dict]:
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        body = f.read()
    if not body.strip():
        return []

    comment_prefix = "#" if language == "python" else "//"
    header = f"{comment_prefix} {language} code from {os.path.basename(filepath)}:"
    text = f"{header}\n\n{body}"
    return [{"text": text, "source": filepath, "doc_type": f"code_{language}"}]


def _extract_binary_text(raw: bytes, min_len: int = 6) -> str:
    printable = []
    current = []
    for b in raw:
        ch = chr(b)
        if 32 <= b <= 126:
            current.append(ch)
        else:
            if len(current) >= min_len:
                printable.append("".join(current))
            current = []
    if len(current) >= min_len:
        printable.append("".join(current))

    # Deduplicate preserving order
    seen = set()
    ordered = []
    for s in printable:
        if s not in seen:
            seen.add(s)
            ordered.append(s)
    return "\n".join(ordered)


def _load_plc_file(filepath: str, ext: str) -> List[Dict]:
    # Quick support layer retained for compatibility; text-first with binary fallback.
    if ext in PLC_XML_EXTENSIONS:
        return load_text(filepath)

    if ext in PLC_ARCHIVE_EXTENSIONS:
        entries = []
        try:
            with zipfile.ZipFile(filepath, "r") as zf:
                for name in zf.namelist():
                    if name.endswith("/"):
                        continue
                    inner_ext = os.path.splitext(name)[1].lower()
                    if inner_ext in {".xml", ".txt", ".csv", ".scl", ".awl", ".stl", ".cfg", ".ini", ".json", ".log", ".db", ".fc", ".fb"}:
                        content = zf.read(name)
                        text = content.decode("utf-8", errors="ignore")
                        if text.strip():
                            entries.append(
                                {
                                    "text": f"Archive file: {name}\n{text}",
                                    "source": filepath,
                                    "doc_type": "plc_archive",
                                }
                            )
            return entries
        except Exception:
            pass

    try:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        if text.strip():
            return [{"text": text, "source": filepath, "doc_type": "plc"}]
    except Exception:
        pass

    with open(filepath, "rb") as f:
        raw = f.read()
    text = _extract_binary_text(raw)
    return [{"text": text, "source": filepath, "doc_type": "plc_binary"}] if text.strip() else []


def load_all_documents(docs_dir=DOCS_DIR):
    """Walk all subfolders and load every supported document into entry dicts."""
    documents: List[Dict] = []

    supported = {
        ".pdf": load_pdf,
        ".docx": load_word,
        ".xlsx": load_excel,
        ".txt": load_text,
        ".csv": load_csv,
    }

    for root, _, files in os.walk(docs_dir):
        for filename in files:
            filepath = os.path.join(root, filename)
            ext = os.path.splitext(filename)[1].lower()
            try:
                if ext in supported:
                    entries = supported[ext](filepath)
                elif ext in CODE_EXTENSIONS:
                    entries = load_code(filepath, CODE_EXTENSIONS[ext])
                elif ext in ALL_PLC_EXTENSIONS:
                    entries = _load_plc_file(filepath, ext)
                else:
                    continue

                if entries:
                    documents.extend(entries)
            except Exception as exc:
                print(f"ERROR loading {filepath}: {exc}")

    return documents


# -----------------------------
# Incremental ingestion
# -----------------------------
def compute_file_hash(filepath: str) -> str:
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_hash_store(path=HASH_STORE_PATH) -> Dict[str, str]:
    if not os.path.isfile(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_hash_store(hash_map: Dict[str, str], path=HASH_STORE_PATH):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(hash_map, f, indent=2, sort_keys=True)


def _delete_source_chunks(collection, source_path: str):
    existing = collection.get(where={"source": source_path})
    ids = existing.get("ids", []) if existing else []
    if ids:
        collection.delete(ids=ids)


def ingest(force: bool = False):
    print("\n=== SAGE Ingestion Pipeline ===\n")

    os.makedirs(CHROMA_DIR, exist_ok=True)
    os.makedirs(DOCS_DIR, exist_ok=True)

    client = chromadb.PersistentClient(path=CHROMA_DIR)
    embedding_fn = SentenceTransformerEmbeddingFunction(model_name=EMBEDDING_MODEL)
    collection = client.get_or_create_collection(name=COLLECTION_NAME, embedding_function=embedding_fn)

    old_hashes = load_hash_store(HASH_STORE_PATH)
    new_hashes: Dict[str, str] = {}

    all_files = []
    for root, _, files in os.walk(DOCS_DIR):
        for filename in files:
            ext = os.path.splitext(filename)[1].lower()
            if ext in {".pdf", ".docx", ".xlsx", ".txt", ".csv"} | set(CODE_EXTENSIONS.keys()) | ALL_PLC_EXTENSIONS:
                all_files.append(os.path.join(root, filename))

    changed_files = []
    for fp in all_files:
        file_hash = compute_file_hash(fp)
        new_hashes[fp] = file_hash
        if force or old_hashes.get(fp) != file_hash:
            changed_files.append(fp)

    deleted_files = set(old_hashes.keys()) - set(new_hashes.keys())
    for fp in deleted_files:
        _delete_source_chunks(collection, fp)

    if not changed_files and not deleted_files and not force:
        print("No changed files detected. Ingestion skipped.")
        return

    chunk_id = 0
    for filepath in changed_files:
        ext = os.path.splitext(filepath)[1].lower()
        if ext == ".pdf":
            entries = load_pdf(filepath)
        elif ext == ".docx":
            entries = load_word(filepath)
        elif ext == ".xlsx":
            entries = load_excel(filepath)
        elif ext == ".txt":
            entries = load_text(filepath)
        elif ext == ".csv":
            entries = load_csv(filepath)
        elif ext in CODE_EXTENSIONS:
            entries = load_code(filepath, CODE_EXTENSIONS[ext])
        elif ext in ALL_PLC_EXTENSIONS:
            entries = _load_plc_file(filepath, ext)
        else:
            entries = []

        _delete_source_chunks(collection, filepath)

        for entry in entries:
            chunks = chunk_text_semantic(entry["text"])
            for idx, chunk in enumerate(chunks):
                if not chunk.strip():
                    continue

                metadata = {
                    "source": entry["source"],
                    "doc_type": entry.get("doc_type", "unknown"),
                }
                if entry.get("page") is not None:
                    metadata["page"] = int(entry["page"])

                stable_id = f"{hashlib.md5((entry['source'] + str(entry.get('page',''))).encode()).hexdigest()}_{idx}_{chunk_id}"
                collection.add(documents=[chunk], metadatas=[metadata], ids=[stable_id])
                chunk_id += 1

    save_hash_store(new_hashes, HASH_STORE_PATH)

    print(f"✅ Ingestion complete. Added/updated {chunk_id} chunks.")
    print(f"Collection: {COLLECTION_NAME}")
    print(f"Chroma path: {CHROMA_DIR}\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="SAGE ingestion pipeline")
    parser.add_argument("--force", action="store_true", help="Force full re-ingestion")
    args = parser.parse_args()
    ingest(force=args.force)
